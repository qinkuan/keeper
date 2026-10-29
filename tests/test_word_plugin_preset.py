"""word 插件「版式预设」的自测。

跑法：keeper/.venv/bin/python keeper/tests/test_word_plugin_edit.py 同目录的
     test_word_plugin_preset.py

覆盖：四套版式（公文/合同/报告/技术文档）的字号、行距、首行缩进、页边距、
标题样式确实落到 docx 上；显式参数覆盖预设；非法 preset 报业务错误。
"""
from __future__ import annotations

import importlib.util
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

PASS = 0
FAIL = 0


def check(label: str, cond: bool, extra: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {label}")
    else:
        FAIL += 1
        print(f"  ✗ {label}{('  → ' + extra) if extra else ''}")


def _load_plugin():
    path = Path(__file__).resolve().parents[1] / "plugins" / "word" / "bin" / "word.py"
    spec = importlib.util.spec_from_file_location("word_plugin_preset", path)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


def _blocks():
    return [
        {"type": "heading", "level": 1, "text": "一、总则"},
        {"type": "paragraph", "text": "这是正文第一段。"},
        {"type": "bullets", "items": ["要点一", "要点二"]},
        {"type": "table", "header": ["列一", "列二"], "rows": [["值一", "值二"]]},
    ]


def main() -> None:
    mod = _load_plugin()
    tmp = Path(tempfile.mkdtemp())
    ctx = {"workspace": str(tmp), "agent_space": str(tmp), "read_only": False}

    from docx import Document
    from docx.shared import Pt

    def build(name: str, preset: str | None, extra: dict | None = None):
        cfg = {"title": "测试标题", "blocks": _blocks()}
        if preset:
            cfg["preset"] = preset
        if extra:
            cfg.update(extra)
        doc = Document()
        sty = mod._resolve_preset(cfg)
        mod._build(doc, cfg, ctx, sty)
        path = tmp / f"{name}.docx"
        doc.save(str(path))
        return path

    print("\n=== 公文版式 ===")
    d = build("gongwen", "公文")
    doc = Document(str(d))
    normal = doc.styles["Normal"]
    pf = normal.paragraph_format
    check("正文字号 = 16pt（三号）", normal.font.size == Pt(16), str(normal.font.size))
    check("固定行距 28 磅", pf.line_spacing == Pt(28), str(pf.line_spacing))
    check("首行缩进 ≈ 1.13cm（两字符）", abs((pf.first_line_indent.cm or 0) - 1.13) < 0.02,
          str(pf.first_line_indent))
    check("段后 0", pf.space_after == Pt(0), str(pf.space_after))
    h1 = doc.styles["Heading 1"]
    check("一级标题 16pt 黑体", h1.font.size == Pt(16), str(h1.font.size))
    sec = doc.sections[0]
    check("页边距：上 3.7 / 左 2.8",
          abs(sec.top_margin.cm - 3.7) < 0.02 and abs(sec.left_margin.cm - 2.8) < 0.02,
          f"{sec.top_margin.cm}/{sec.left_margin.cm}")
    # 列表不该带首行缩进
    lp = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    check("列表段落首行缩进归零",
          bool(lp) and (lp[0].paragraph_format.first_line_indent is None
                        or lp[0].paragraph_format.first_line_indent.cm == 0),
          str(lp[0].paragraph_format.first_line_indent) if lp else "无列表段")
    # 表格单元格也不该缩进
    cell_p = doc.tables[0].cell(0, 0).paragraphs[0]
    check("表格单元格首行缩进归零",
          cell_p.paragraph_format.first_line_indent is None
          or cell_p.paragraph_format.first_line_indent.cm == 0,
          str(cell_p.paragraph_format.first_line_indent))

    print("\n=== 合同版式 ===")
    d = build("contract", "contract")
    doc = Document(str(d))
    normal = doc.styles["Normal"]
    check("正文字号 = 12pt", normal.font.size == Pt(12), str(normal.font.size))
    check("行距 1.5 倍", abs(float(normal.paragraph_format.line_spacing) - 1.5) < 0.01,
          str(normal.paragraph_format.line_spacing))
    check("首行缩进 ≈ 0.74cm", abs((normal.paragraph_format.first_line_indent.cm or 0) - 0.74) < 0.02,
          str(normal.paragraph_format.first_line_indent))
    sec = doc.sections[0]
    check("页边距 2.54 / 3.17",
          abs(sec.top_margin.cm - 2.54) < 0.02 and abs(sec.left_margin.cm - 3.17) < 0.02,
          f"{sec.top_margin.cm}/{sec.left_margin.cm}")

    print("\n=== 报告版式（默认带目录） ===")
    d = build("report", "报告")
    doc = Document(str(d))
    xml = doc.element.body.xml
    check("目录域已插入", "TOC" in xml, "没找到 TOC 域")
    check("正文字号 = 12pt", doc.styles["Normal"].font.size == Pt(12))

    print("\n=== 技术文档版式（不首行缩进、带目录） ===")
    d = build("tech", "tech")
    doc = Document(str(d))
    normal = doc.styles["Normal"]
    check("正文字号 = 11pt", normal.font.size == Pt(11), str(normal.font.size))
    check("没有首行缩进",
          normal.paragraph_format.first_line_indent is None
          or (normal.paragraph_format.first_line_indent.cm or 0) == 0,
          str(normal.paragraph_format.first_line_indent))
    check("默认带目录（说明书也要目录）", "TOC" in doc.element.body.xml)
    check("代码块字号 10pt（技术版式）", mod.PRESETS["tech"]["code_size"] == 10)

    print("\n=== 显式值覆盖预设 ===")
    d = build("override", "gongwen", {"font_ea": "宋体", "page": {"size": "A4", "margin_cm": 2.0}})
    doc = Document(str(d))
    check("显式 font_ea 覆盖预设", "宋体" in (doc.styles["Normal"].element.xml),
          "Normal 里没找到宋体")
    sec = doc.sections[0]
    check("显式 margin_cm 覆盖预设", abs(sec.top_margin.cm - 2.0) < 0.02, str(sec.top_margin.cm))

    print("\n=== 非法 preset 报错 ===")
    try:
        build("bad", "小清新")
        check("非法 preset 被拒", False, "居然没报错")
    except mod.Biz as e:
        check("非法 preset 被拒", "preset 无效" in str(e), str(e))

    print("\n=== 不给 preset = 保持旧行为 ===")
    d = build("plain", None)
    doc = Document(str(d))
    check("默认正文 12pt / 1.5 倍行距",
          doc.styles["Normal"].font.size == Pt(12)
          and abs(float(doc.styles["Normal"].paragraph_format.line_spacing) - 1.5) < 0.01)

    print(f"\n结果：{PASS} 通过 / {FAIL} 失败")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
