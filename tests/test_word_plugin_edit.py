"""word 插件「受限回写」（edit）的自测。

跑法：keeper/.venv/bin/python keeper/tests/test_word_plugin_edit.py

覆盖：追加段 / 替换段文本 / 删段；**原件不被覆盖**；表格与格式不丢；
越界与非法参数被当成业务错误（而不是崩）。
"""
from __future__ import annotations

import importlib.util
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

PASS = 0
FAIL = 0


def check(label: str, cond: bool, extra: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {label}")
    else:
        FAIL += 1
        print(f"  ✗ {label}{('  → ' + extra) if extra else ''}")


def _load_plugin():
    path = Path(__file__).resolve().parents[1] / "plugins" / "word" / "bin" / "word.py"
    spec = importlib.util.spec_from_file_location("word_plugin", path)
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


def _make_docx(p: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段：原始内容")
    doc.add_paragraph("第二段：待替换")
    doc.add_paragraph("第三段：待删除")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "表格格一"
    t.cell(0, 1).text = "表格格二"
    doc.add_paragraph("第四段：文末")
    doc.save(str(p))


def main() -> None:
    mod = _load_plugin()
    tmp = Path(tempfile.mkdtemp())
    src = tmp / "原文档.docx"
    _make_docx(src)
    before_bytes = src.read_bytes()

    ctx = {"workspace": str(tmp), "agent_space": str(tmp), "read_only": False}

    print("\n=== 追加 / 替换 / 删除 ===")
    res = mod.handle(
        "edit",
        {
            "path": "原文档.docx",
            "out_path": "改后.docx",
            "ops": [
                {"op": "replace", "index": 2, "text": "第二段：已改写"},
                {"op": "delete", "index": 3},
                {"op": "append", "after": 1, "text": "新插入的一段"},
                {"op": "append", "text": "文末追加的一段"},
            ],
        },
        ctx,
    )
    print("  返回：", res.replace("\n", " / "))

    out = tmp / "改后.docx"
    check("副本已生成", out.exists())
    check("原件字节未变", src.read_bytes() == before_bytes)

    from docx import Document

    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    check(
        "段落顺序正确",
        texts
        == [
            "第一段：原始内容",
            "新插入的一段",
            "第二段：已改写",
            "第四段：文末",
            "文末追加的一段",
        ],
        str(texts),
    )
    check(
        "表格还在（内容没丢）",
        len(doc.tables) == 1 and doc.tables[0].cell(0, 0).text == "表格格一",
        str(len(doc.tables)),
    )

    print("\n=== 多 run 的段落替换 ===")
    doc2 = Document()
    p = doc2.add_paragraph()
    r1 = p.add_run("前半")
    r1.bold = True
    p.add_run("后半")
    src2 = tmp / "多run.docx"
    doc2.save(str(src2))
    mod.handle(
        "edit",
        {
            "path": "多run.docx",
            "out_path": "多run_改.docx",
            "ops": [{"op": "replace", "index": 1, "text": "整段换掉了"}],
        },
        ctx,
    )
    d3 = Document(str(tmp / "多run_改.docx"))
    check("整段文本被替换", d3.paragraphs[0].text == "整段换掉了", d3.paragraphs[0].text)
    check(
        "首个 run 的格式保留（加粗）",
        bool(d3.paragraphs[0].runs) and bool(d3.paragraphs[0].runs[0].bold),
        str([(r.text, r.bold) for r in d3.paragraphs[0].runs]),
    )

    print("\n=== 拒绝覆盖原件 / 越界 / 非法参数 ===")
    try:
        mod.handle(
            "edit",
            {
                "path": "原文档.docx",
                "out_path": "原文档.docx",
                "ops": [{"op": "delete", "index": 1}],
            },
            ctx,
        )
        check("拒绝覆盖原件", False, "居然没报错")
    except mod.Biz as e:
        check("拒绝覆盖原件", "覆盖原件" in str(e), str(e))

    try:
        mod.handle(
            "edit",
            {
                "path": "原文档.docx",
                "out_path": "x.docx",
                "ops": [{"op": "replace", "index": 99, "text": "x"}],
            },
            ctx,
        )
        check("段号越界被拒", False, "居然没报错")
    except mod.Biz as e:
        check("段号越界被拒", "越界" in str(e), str(e))

    try:
        mod.handle(
            "edit",
            {
                "path": "原文档.docx",
                "out_path": "x.docx",
                "ops": [{"op": "rename_everything"}],
            },
            ctx,
        )
        check("未知操作被拒", False, "居然没报错")
    except mod.Biz as e:
        check("未知操作被拒", "未知操作" in str(e), str(e))

    print("\n=== 路径越界 ===")
    try:
        mod.handle(
            "edit",
            {"path": "../别处.docx", "out_path": "y.docx",
             "ops": [{"op": "delete", "index": 1}]},
            ctx,
        )
        check("工作区外的路径被拒", False, "居然没报错")
    except mod.Biz as e:
        check("工作区外的路径被拒", "越界" in str(e), str(e))

    print(f"\n结果：{PASS} 通过 / {FAIL} 失败")
    sys.exit(1 if FAIL else 0)


if __name__ == "__main__":
    main()
